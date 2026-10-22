#!/usr/bin/env python3
"""Verify the formula content of a structured niche.

1. Lint every formula (balanced parentheses / quotes, starts with '=').
2. Recalculate every worked example with LibreOffice (headless) and compare each formula
   result with the result printed on the page. Excel formulas are checked; Google Sheets
   formulas are checked too when they differ and use only functions LibreOffice knows.
   Checks marked `verify: false` (functions LibreOffice lacks, e.g. TEXTSPLIT, REGEXEXTRACT)
   are reported as "not machine-verified".

Usage: python scripts/verify_formulas.py [--niche NAME] [--only ID] [--keep]
Exit 0 = all verified checks match. Requires `soffice` (LibreOffice) on PATH.
"""
import argparse
import datetime as dt
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import structured
from common import load_niche, load_site

XLFN = ["XLOOKUP", "XMATCH", "UNIQUE", "SEQUENCE", "TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS",
        "MINIFS", "IFNA", "DAYS", "TEXTBEFORE", "TEXTAFTER", "TEXTSPLIT", "CHOOSECOLS", "CHOOSEROWS",
        "HSTACK", "VSTACK", "TAKE", "DROP", "RRI", "STDEV.S", "RANK.EQ", "LET", "ISOWEEKNUM",
        "NETWORKDAYS.INTL", "WORKDAY.INTL", "AGGREGATE", "FORMULATEXT", "TOCOL"]
XLWS = ["FILTER", "SORT"]
SHEETS_ONLY = {"ARRAYFORMULA", "SPLIT", "REGEXEXTRACT", "REGEXMATCH", "REGEXREPLACE", "QUERY",
               "IMPORTRANGE", "JOIN", "COUNTUNIQUE", "AVERAGE.WEIGHTED", "GOOGLEFINANCE", "FLATTEN"}
# Functions LibreOffice 25.2 cannot calculate: formulas using them are linted but not recalculated.
LO_UNSUPPORTED = {"CHOOSECOLS", "CHOOSEROWS", "HSTACK", "VSTACK", "TAKE", "DROP", "TEXTSPLIT",
                  "TEXTBEFORE", "TEXTAFTER", "TOCOL", "WORKDAY.INTL", "ISOWEEKNUM"}
EPOCH = dt.date(1899, 12, 30)
UNVERIFIED = []


def lint(f):
    errs = []
    if not f.strip().startswith("="):
        errs.append("does not start with =")
    depth, q, b = 0, False, 0
    for ch in f:
        if ch == '"':
            q = not q
        elif not q:
            depth += ch == "("
            depth -= ch == ")"
            b += ch == "{"
            b -= ch == "}"
            if depth < 0:
                errs.append("unbalanced )")
                break
    if q:
        errs.append("unbalanced quotes")
    if depth > 0:
        errs.append("unclosed (")
    if b:
        errs.append("unbalanced {}")
    return errs


def funcs(f):
    return set(re.findall(r"(?<![A-Za-z0-9_.\"])([A-Z][A-Z0-9.]*)\(", re.sub(r'"[^"]*"', '""', f)))


def to_xlsx_formula(f):
    f = f.strip()
    out, parts = [], re.split(r'("[^"]*")', f)
    for part in parts:
        if part.startswith('"'):
            out.append(part)
            continue
        for fn in XLWS:
            part = re.sub(rf"(?<![A-Za-z0-9_.]){fn}\(", f"_xlfn._xlws.{fn}(", part)
        for fn in XLFN:
            part = re.sub(rf"(?<![A-Za-z0-9_.]){re.escape(fn)}\(", f"_xlfn.{fn}(", part)
        out.append(part)
    return "".join(out)


REF = re.compile(r"(?<![A-Za-z0-9_.!'\"])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_])")


def shift(f, dr, dc):
    def col(s):
        n = 0
        for ch in s:
            n = n * 26 + ord(ch) - 64
        return n

    def rep(m):
        c = m.group(2) if m.group(1) else structured.col_letter(col(m.group(2)) - 1 + dc)
        r = m.group(4) if m.group(3) else str(int(m.group(4)) + dr)
        return f"{m.group(1)}{c}{m.group(3)}{r}"
    parts = re.split(r'("[^"]*")', f)
    return "".join(p if p.startswith('"') else REF.sub(rep, p) for p in parts)


def cell_value(v):
    if isinstance(v, str):
        if re.fullmatch(r"\d{1,3}:\d{2}", v):
            h, m = v.split(":")
            return int(h) / 24 + int(m) / 1440
        if re.fullmatch(r"-?\d+(\.\d+)?%", v):
            return float(v[:-1]) / 100
    return v


def expect_value(v):
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        return float((v - EPOCH).days)
    v = cell_value(v)
    if v is None:
        return ""
    return v


def same(got, exp):
    if isinstance(got, dt.datetime):
        got = (got - dt.datetime(1899, 12, 30)).total_seconds() / 86400
    if isinstance(got, dt.time):
        got = (got.hour * 3600 + got.minute * 60 + got.second) / 86400
    if got is None:
        got = ""
    if isinstance(exp, bool) or isinstance(got, bool):
        return got == exp
    if isinstance(exp, (int, float)) and isinstance(got, (int, float)):
        return abs(got - exp) <= max(0.0051, abs(exp) * 1e-9)
    if isinstance(exp, (int, float)) and isinstance(got, str):
        try:
            return abs(float(got) - exp) <= 0.0051
        except ValueError:
            return False
    return str(got) == str(exp)


def plan_checks(e):
    """Yield (platform, check, cell, formula, kind, ref_range, expected)."""
    for c in e["example"].get("checks", []):
        if c.get("verify") is False:
            continue
        variants = [("excel", c.get("excel") or c.get("formula"))]
        sf = c.get("sheets")
        if sf and sf != variants[0][1]:
            variants.append(("sheets", sf))
        for plat, f in variants:
            if not f:
                continue
            bad = funcs(f) & (LO_UNSUPPORTED | SHEETS_ONLY)
            if bad:
                UNVERIFIED.append(f"{e['id']} {c['cell']} ({plat}): uses {', '.join(sorted(bad))}")
                continue
            if c.get("fill"):
                r0, c0 = structured.cell_rc(c["cell"])
                for ref, exp in zip(c["fill"], c.get("expects") or c["results"]):
                    r1, c1 = structured.cell_rc(ref)
                    yield plat, c, ref, shift(f, r1 - r0, c1 - c0), "cell", None, exp
            elif c.get("spill_cells"):
                yield plat, c, c["cell"], f, "spill", c["spill_cells"], c["spill_results"]
            else:
                yield plat, c, c["cell"], f, "array" if c.get("array") else "cell", None, c["expect"] if "expect" in c else c.get("result")


def verify_entries(entries, keep=False):
    """Lint + recalculate. Returns (problems, n_matched, unverified)."""
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
    UNVERIFIED.clear()
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return ["soffice (LibreOffice) not found"], 0, []
    problems, unverified = [], []
    for e in entries:
        fs = []
        for p in ("excel", "sheets"):
            b = e.get(p) or {}
            fs += [b.get("formula")] + [x.get("formula") for x in b.get("alternatives") or []]
        for c in (e.get("example") or {}).get("checks", []):
            fs += [c.get("excel"), c.get("sheets"), c.get("formula")]
        for f in filter(None, fs):
            for err in lint(f):
                problems.append(f"{e['id']}: lint: {err}: {f}")
        for c in (e.get("example") or {}).get("checks", []):
            if c.get("verify") is False:
                unverified.append(f"{e['id']} {c['cell']}: {c.get('why_unverified', 'not supported by LibreOffice')}")

    tmp = Path(tempfile.mkdtemp(prefix="verify-"))
    jobs = {}
    for e in entries:
        for plat in ("excel", "sheets"):
            try:
                plan = [x for x in plan_checks(e) if x[0] == plat]
            except Exception as ex:  # malformed check data
                problems.append(f"{e['id']}: bad check data: {ex}")
                break
            if not plan:
                continue
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Main"
            for r, row in enumerate(e["example"]["grid"], 1):
                for ci, v in enumerate(row, 1):
                    if v is not None:
                        ws.cell(r, ci, cell_value(v))
            for name, og in (e["example"].get("other_sheets") or {}).items():
                w2 = wb.create_sheet(name)
                for r, row in enumerate(og, 1):
                    for ci, v in enumerate(row, 1):
                        if v is not None:
                            w2.cell(r, ci, cell_value(v))
            for _, c, ref, f, kind, rng, exp in plan:
                xf = to_xlsx_formula(f)
                if kind == "spill":
                    ws[ref] = ArrayFormula(f"{rng[0]}:{rng[-1]}", xf)
                elif kind == "array":
                    ws[ref] = ArrayFormula(f"{ref}:{ref}", xf)
                else:
                    ws[ref] = xf
            fn = tmp / f"{re.sub(r'[^A-Za-z0-9_-]', '_', e['id'])}__{plat}.xlsx"
            wb.save(fn)
            jobs[fn.name] = (e, plat, plan)
    if jobs:
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "out")]
                       + [str(tmp / n) for n in jobs], capture_output=True, timeout=600)
    n_ok = 0
    for name, (e, plat, plan) in jobs.items():
        f = tmp / "out" / name
        if not f.exists():
            problems.append(f"{e['id']} ({plat}): LibreOffice produced no output")
            continue
        ws = openpyxl.load_workbook(f, data_only=True)["Main"]
        for _, c, ref, formula, kind, rng, exp in plan:
            pairs = list(zip(rng, exp)) if kind == "spill" else [(ref, exp)]
            for cref, ev in pairs:
                got = ws[cref].value
                if same(got, expect_value(ev)):
                    n_ok += 1
                else:
                    problems.append(f"{e['id']} ({plat}) {cref}: expected {ev!r}, LibreOffice got {got!r}  [{formula}]")
    if keep:
        print(f"[verify] workbooks kept in {tmp}")
    else:
        shutil.rmtree(tmp, ignore_errors=True)
    return problems, n_ok, unverified + sorted(set(UNVERIFIED))


def main():
    site = load_site()
    ap = argparse.ArgumentParser()
    ap.add_argument("--niche", default=site["niche"])
    ap.add_argument("--only")
    ap.add_argument("--keep", action="store_true")
    a = ap.parse_args()
    niche = load_niche(a.niche)
    entries = [e for e in structured.load_entries(niche) if not a.only or e["id"] == a.only]
    problems, n_ok, unverified = verify_entries(entries, a.keep)
    for p in problems:
        print("[verify] PROBLEM", p)
    for u in unverified:
        print("[verify] not machine-verified:", u)
    print(f"[verify] entries={len(entries)} results matched={n_ok} problems={len(problems)} "
          f"unverified checks={len(unverified)}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
